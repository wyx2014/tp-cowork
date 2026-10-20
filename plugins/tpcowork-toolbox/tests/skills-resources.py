#!/usr/bin/env python3
"""Run retained skill resources and verify the packaged document libraries."""
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile

from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[3]
SKILLS = ROOT / "plugins/tpcowork-toolbox/resources/skills"
ARTIFACTS = ROOT / "artifacts/skills-validation"
ARTIFACTS.mkdir(parents=True, exist_ok=True)
run = Path(tempfile.mkdtemp(prefix="resources-", dir=ARTIFACTS))
env = {key: value for key, value in os.environ.items() if not any(word in key.upper() for word in ("KEY", "SECRET", "TOKEN", "PASSWORD", "CREDENTIAL"))}
env["PYTHONDONTWRITEBYTECODE"] = "1"


def execute(script, *args):
    result = subprocess.run([sys.executable, str(script), *map(str, args)], capture_output=True, text=True, env=env, timeout=60)
    if result.returncode:
        raise RuntimeError(f"{script.name}: {result.stderr or result.stdout}")
    return result.stdout


pptx = run / "migration-sample.pptx"
sample = Presentation()
sample.slides.add_slide(sample.slide_layouts[0]).shapes.title.text = "文档运行库验证"
sample.slides.add_slide(sample.slide_layouts[1]).shapes.title.text = "可编辑指标"
sample.save(pptx)
presentation = Presentation(pptx)
assert len(presentation.slides) == 2
assert any("可编辑指标" in shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)

docx = run / "migration-sample.docx"
sample = Document()
sample.add_heading("资源迁移验证", level=1)
sample.add_paragraph("本文为测试样例，不含用户业务数据。")
sample.add_heading("待办", level=2)
sample.add_paragraph("负责人：待确认")
sample.save(docx)
document = Document(docx)
assert [p.text for p in document.paragraphs] == ["资源迁移验证", "本文为测试样例，不含用户业务数据。", "待办", "负责人：待确认"]
for file in (pptx, docx):
    with zipfile.ZipFile(file) as archive:
        assert archive.testzip() is None

creator = SKILLS / "skill-creator/scripts"
execute(creator / "quick_validate.py", SKILLS / "skill-creator")
execute(creator / "package_skill.py", SKILLS / "skill-creator", run)
with zipfile.ZipFile(run / "skill-creator.skill") as archive:
    assert any(name.endswith("scripts/init_skill.py") for name in archive.namelist())
    assert any(name.endswith("references/gui-authoring.md") for name in archive.namelist())

report = {"pptx": {"slides": len(presentation.slides), "source": "bundled python-pptx module"}, "docx": {"paragraphs": len(document.paragraphs), "source": "bundled python-docx module"}, "skillCreator": "validated and packaged", "dependencies": {name: importlib.util.find_spec(name) is not None for name in ("pptx", "PIL", "yaml", "docx")}}
(run / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"result": "passed", "artifacts": str(run.relative_to(ROOT)), "slides": len(presentation.slides), "docxParagraphs": 4}))
